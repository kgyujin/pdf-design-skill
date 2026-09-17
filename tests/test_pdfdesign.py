"""
작성자: Git 이력 참조
작성목적: pdfdesign.py 의 디자인 설정 로직과 CSS 축 정합성을 검증한다.
작성일: 2026-09-17
주의사항:
  - 표준 라이브러리 unittest 만 사용한다. 실행: python3 -m unittest discover -s tests
  - 렌더링 테스트는 Chrome과 pdfinfo가 있을 때만 실행하고, 없으면 skip 한다.
변경사항 내역:
- 2026-09-17 | 최초 작성 | set 부분 변경, accent, prefs, CSS 축 정합성, 렌더 스모크
- 2026-09-17 | 발표 자료 모드 | 노트 파서, 시간 추정, DOCX 구조, check 결함 탐지, talk 렌더·대본
- 2026-09-17 | 바탕체 규칙 | 바탕체는 font 옵션에서만, 폰트는 고운바탕
- 2026-09-17 | AI 티 검사 | audit_design 탐지·오탐 방지, 기본 CSS·템플릿에 AI 패턴이 없는지 검사
"""

from __future__ import annotations

import contextlib
import io
import json
import re
import shutil
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.dom import minidom

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT / "scripts"))

import deck_tools  # noqa: E402
import pdfdesign  # noqa: E402
from docx_writer import DocxBuilder  # noqa: E402

HAS_RENDERER = bool(shutil.which("pdfinfo") and shutil.which("pdftoppm")) and any(
    Path(path).exists() or shutil.which(path) for path in pdfdesign.CHROME_CANDIDATES
)
TINY_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000002000000010806000000f478d4fa"
    "0000000d49444154789c63f8cf00020c0c0001050101e1d6d1f70000000049454e44ae426082"
)


def run_cli(*argv: str) -> tuple[int, str, str]:
    stdout, stderr = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(stdout), contextlib.redirect_stderr(stderr):
        exit_code = pdfdesign.main(list(argv))
    return exit_code, stdout.getvalue(), stderr.getvalue()


class DesignSettingTest(unittest.TestCase):
    def setUp(self) -> None:
        self.work_dir = Path(tempfile.mkdtemp(prefix="pdf-design-test-"))
        self.doc = self.work_dir / "doc.html"
        exit_code, _, stderr = run_cli("init", "report", str(self.doc), "--style", "bold", "--palette", "forest")
        self.assertEqual(exit_code, 0, stderr)

    def tearDown(self) -> None:
        shutil.rmtree(self.work_dir, ignore_errors=True)

    def design(self) -> dict[str, str]:
        return pdfdesign.read_design(self.doc.read_text(encoding="utf-8"))

    def test_init_applies_options_and_writes_bundle(self) -> None:
        self.assertEqual(self.design()["style"], "bold")
        self.assertEqual(self.design()["palette"], "forest")
        self.assertTrue((self.work_dir / pdfdesign.BUNDLE_NAME).exists())

    def test_set_changes_only_requested_axis(self) -> None:
        run_cli("set", str(self.doc), "--style", "editorial")
        design = self.design()
        self.assertEqual(design["style"], "editorial")
        self.assertEqual(design["palette"], "forest", "색조는 유지돼야 한다")
        self.assertEqual(design["density"], "normal")

    def test_set_keeps_document_content(self) -> None:
        before_body = self.doc.read_text(encoding="utf-8").split("</head>", 1)[1]
        run_cli("set", str(self.doc), "--palette", "plum", "--density", "compact")
        after_body = self.doc.read_text(encoding="utf-8").split("</head>", 1)[1]
        self.assertEqual(before_body, after_body)

    def test_accent_is_added_replaced_and_removed(self) -> None:
        run_cli("set", str(self.doc), "--accent", "#FF6600")
        self.assertEqual(self.design()["accent"], "#FF6600")
        run_cli("set", str(self.doc), "--accent", "#123456")
        text = self.doc.read_text(encoding="utf-8")
        self.assertEqual(text.count(pdfdesign.ACCENT_STYLE_ID), 1, "포인트색 블록은 하나만 있어야 한다")
        self.assertEqual(self.design()["accent"], "#123456")
        self.assertEqual(self.design()["palette"], "forest")
        run_cli("set", str(self.doc), "--accent", "none")
        self.assertNotIn(pdfdesign.ACCENT_STYLE_ID, self.doc.read_text(encoding="utf-8"))

    def test_invalid_accent_is_rejected_without_writing(self) -> None:
        original = self.doc.read_text(encoding="utf-8")
        exit_code, _, stderr = run_cli("set", str(self.doc), "--accent", "orange")
        self.assertEqual(exit_code, 1)
        self.assertIn("#RRGGBB", stderr)
        self.assertEqual(self.doc.read_text(encoding="utf-8"), original)

    def test_saved_prefs_become_defaults_for_next_init(self) -> None:
        run_cli("set", str(self.doc), "--radius", "round", "--save")
        prefs = json.loads((self.work_dir / pdfdesign.PREFS_NAME).read_text(encoding="utf-8"))
        self.assertEqual(prefs["radius"], "round")

        next_doc = self.work_dir / "deck.html"
        run_cli("init", "deck", str(next_doc), "--palette", "navy")
        next_design = pdfdesign.read_design(next_doc.read_text(encoding="utf-8"))
        self.assertEqual(next_design["style"], "bold", "저장된 기본값을 따라야 한다")
        self.assertEqual(next_design["radius"], "round")
        self.assertEqual(next_design["palette"], "navy", "CLI 옵션이 저장된 기본값보다 우선한다")

    def test_init_refuses_to_overwrite(self) -> None:
        exit_code, _, stderr = run_cli("init", "report", str(self.doc))
        self.assertEqual(exit_code, 1)
        self.assertIn("--force", stderr)


class GalleryOptionTest(unittest.TestCase):
    def test_parse_axis_values(self) -> None:
        self.assertEqual(pdfdesign.parse_axis_values("style", None, "tech"), ["tech"])
        self.assertEqual(pdfdesign.parse_axis_values("style", "all", "tech"), list(pdfdesign.DESIGN_AXES["style"]))
        self.assertEqual(pdfdesign.parse_axis_values("palette", "keep,navy,navy", "plum"), ["plum", "navy"])
        with self.assertRaises(pdfdesign.PdfDesignError):
            pdfdesign.parse_axis_values("palette", "rainbow", "plum")

    def test_variant_labels_are_unique(self) -> None:
        labels = [pdfdesign.variant_label(index) for index in range(pdfdesign.GALLERY_MAX_VARIANTS)]
        self.assertEqual(len(labels), len(set(labels)))
        self.assertEqual(labels[26], "AA")


class CssAxisConsistencyTest(unittest.TestCase):
    """DESIGN_AXES 에 있는 값마다 CSS 규칙이 있어야 선택해도 모양이 바뀐다."""

    PALETTE_TOKENS = (
        "--ink", "--ink-2", "--muted", "--line", "--surface", "--surface-2", "--accent", "--accent-soft",
        "--on-accent", "--inverse", "--on-inverse", "--on-inverse-muted", "--accent-on-inverse",
    )
    NEUTRAL_VALUES = {"font": {"auto"}, "density": {"normal"}, "radius": {"auto"}}

    def setUp(self) -> None:
        self.bundle = pdfdesign.build_css_bundle()

    def test_every_axis_value_has_selector(self) -> None:
        for axis, choices in pdfdesign.DESIGN_AXES.items():
            for value in choices:
                if value in self.NEUTRAL_VALUES.get(axis, set()):
                    continue
                self.assertIn(f'data-{axis}="{value}"', self.bundle, f"{axis}={value} 규칙 없음")

    def test_every_style_file_exists(self) -> None:
        for source in pdfdesign.css_sources():
            self.assertTrue(source.exists(), source)

    def test_every_palette_defines_all_tokens(self) -> None:
        palettes_css = (pdfdesign.CSS_DIR / "palettes.css").read_text(encoding="utf-8")
        for name in pdfdesign.DESIGN_AXES["palette"]:
            block = re.search(rf'html\[data-palette="{name}"\]\s*\{{(.*?)\}}', palettes_css, re.DOTALL)
            self.assertIsNotNone(block, name)
            defined = set(re.findall(r"(--[a-z0-9-]+):", block.group(1)))
            self.assertEqual(defined, set(self.PALETTE_TOKENS), f"{name} 팔레트 토큰 불일치")

    def test_serif_is_opt_in_and_uses_gowun_batang(self) -> None:
        # 바탕(명조)체는 사용자가 요청할 때만 font 옵션으로 켠다. 스타일·템플릿·예제가 직접 쓰면 안 된다.
        self.assertIn('--font-serif: "Gowun Batang"', self.bundle)
        self.assertIn("family=Gowun+Batang", self.bundle)
        self.assertNotIn("Noto Serif", self.bundle)
        # 기본 폰트는 Pretendard. 스타일 파일은 세리프·모노 폰트를 직접 지정하지 않는다(옵션 축이 담당).
        allowed = {pdfdesign.CSS_DIR / "core.css", pdfdesign.CSS_DIR / "options.css"}
        for source in pdfdesign.css_sources():
            if source not in allowed:
                style_css = source.read_text(encoding="utf-8")
                self.assertNotIn("--font-serif", style_css, source.name)
                self.assertNotIn("--font-mono", style_css, source.name)
        self.assertIn('--font-sans: "Pretendard"', self.bundle)
        for html_path in [*pdfdesign.TEMPLATE_DIR.glob("*.html"), *(REPO_ROOT / "examples").glob("*.html")]:
            text = html_path.read_text(encoding="utf-8")
            self.assertNotIn("--font-serif", text, html_path.name)
            self.assertNotRegex(text, r'data-font="serif', html_path.name)
        options_css = (pdfdesign.CSS_DIR / "options.css").read_text(encoding="utf-8")
        for value in ("serif", "serif-all"):
            rule = re.search(rf'html\[data-font="{value}"\] \{{([^}}]*)\}}', options_css).group(1)
            self.assertIn("--head-weight: 700", rule, "고운바탕은 700까지만 있어 제목 굵기를 고정해야 한다")

    def test_import_rules_stay_at_top_of_bundle(self) -> None:
        # @import 는 주석 외 다른 규칙보다 앞에 있어야 브라우저가 무시하지 않는다.
        without_comments = re.sub(r"/\*.*?\*/", "", self.bundle, flags=re.DOTALL).lstrip()
        self.assertTrue(without_comments.startswith("@import"))


SAMPLE_DECK = """
<html><head><title>Deck</title><meta name="pdf-design:duration" content="2"></head><body class="deck">
<section class="slide hero">
  <div class="slide-title">Opening</div>
  <footer class="slide-foot"><span class="source"></span><span class="page"></span></footer>
  <aside class="notes"><div class="talk"><p>안녕하세요. 오늘 발표를 시작하겠습니다.</p></div></aside>
</section>
<section class="slide">
  <header class="slide-head"><div class="eyebrow">Result</div><div class="slide-title">Latency <br>dropped 42%</div>
  <p class="slide-msg">Key message here.</p></header>
  <div class="slide-body l-points"><ol class="points"><li><strong>A</strong> text</li></ol></div>
  <footer class="slide-foot"><span class="source">Source: Report, Fig. 1</span><span class="page"></span></footer>
  <aside class="notes">
    <h4>배경지식</h4>
    <p><strong>p95</strong>는 느린 쪽 5% 경계입니다.</p>
    <ul><li>첫째</li><li>둘째<ul><li>세부</li></ul></li></ul>
    <div class="talk"><p>세로축은 응답시간입니다.</p><p>두 번째 문단입니다.</p></div>
  </aside>
</section>
<section class="slide">
  <header class="slide-head"><div class="slide-title">No notes here</div></header>
  <footer class="slide-foot"><span class="source"></span><span class="page"></span></footer>
</section>
</body></html>
"""


class DeckToolsTest(unittest.TestCase):
    def setUp(self) -> None:
        self.slides = deck_tools.parse_slides(SAMPLE_DECK)

    def test_parses_slide_fields(self) -> None:
        self.assertEqual(len(self.slides), 3)
        second = self.slides[1]
        self.assertEqual(second.title, "Latency dropped 42%")
        self.assertEqual(second.message, "Key message here.")
        self.assertEqual(second.source, "Source: Report, Fig. 1")
        self.assertEqual(self.slides[0].kinds, ["hero"])

    def test_notes_blocks_keep_kind_bold_and_nesting(self) -> None:
        blocks = self.slides[1].notes
        self.assertEqual([block.kind for block in blocks], ["heading", "paragraph", "bullet", "bullet", "bullet", "talk", "talk"])
        self.assertEqual(blocks[1].runs[0], ("p95", True))
        self.assertEqual(blocks[4].level, 1, "중첩 목록은 level 1")
        self.assertEqual(self.slides[1].talk_text, "세로축은 응답시간입니다. 두 번째 문단입니다.")

    def test_audit_flags_missing_notes_and_source_but_not_on_hero(self) -> None:
        issues = deck_tools.audit_slides(self.slides)
        where_messages = [(issue["where"], issue["message"]) for issue in issues]
        self.assertTrue(any(where == "slide 3" and "노트" in message for where, message in where_messages))
        self.assertTrue(any(where == "slide 3" and "출처" in message for where, message in where_messages))
        self.assertFalse(any(where in ("slide 1", "slide 2") for where, _ in where_messages))

    def test_estimate_seconds(self) -> None:
        korean = "가" * deck_tools.KOREAN_CHARS_PER_MINUTE
        self.assertEqual(deck_tools.estimate_seconds(korean), 60)
        english = " ".join(["word"] * deck_tools.ENGLISH_WORDS_PER_MINUTE * 2)
        self.assertEqual(deck_tools.estimate_seconds(english), 120)
        self.assertEqual(deck_tools.estimate_seconds("  "), 0)
        self.assertEqual(deck_tools.read_target_minutes(SAMPLE_DECK), 2.0)
        self.assertEqual(deck_tools.format_duration(75), "1분 15초")

    def test_check_script_injected_before_body_end(self) -> None:
        injected = deck_tools.inject_check_script(SAMPLE_DECK)
        self.assertLess(injected.index("pdf-design-check-runner"), injected.rindex("</body>"))
        fake_dom = '<html><body><script type="application/json" id="pdf-design-check-result">{"issues": [], "slides": 3}</script></body></html>'
        self.assertEqual(deck_tools.extract_check_result(fake_dom), {"issues": [], "slides": 3})


AI_TELL_DOC = """<html><head><style>
  .hero { background: linear-gradient(90deg, #6366f1, #a855f7); box-shadow: 0 4px 20px rgba(0,0,0,.2); }
  .label { text-transform: uppercase; letter-spacing: 0.2em; }
  .card { border-left: 4px solid var(--accent); border-radius: 999px; }
</style></head><body>
<h1>🚀 Launch plan</h1>
<p>We move fast — and we ship often.</p>
<p>Quality matters — every single time.</p>
<h2><span class="sec-num">01</span>Intro</h2>
<div class="eyebrow">A</div><div class="eyebrow">B</div><div class="eyebrow">C</div>
<div class="grid-3"><div class="card">1</div><div class="card">2</div><div class="card">3</div></div>
</body></html>"""


class AiTellAuditTest(unittest.TestCase):
    def test_detects_common_patterns(self) -> None:
        messages = " / ".join(issue["message"] for issue in deck_tools.audit_design(AI_TELL_DOC))
        for expected in ("그라데이션", "그림자", "대문자 라벨", "넓은 자간", "왼쪽 굵은 세로 바", "알약", "HEX", "이모지", "줄표", "0을 붙인 번호", "라벨이 3개", "카드 3개"):
            self.assertIn(expected, messages)
        self.assertTrue(all(issue["message"].startswith(deck_tools.AI_TELL_PREFIX) for issue in deck_tools.audit_design(AI_TELL_DOC)))

    def test_ignores_functional_markup(self) -> None:
        clean = """<html><head><style>@page { size: A4; } #face { color: var(--ink); }</style>
        <style id="pdf-design-accent">html[data-palette] { --accent: #123456; }</style></head><body>
        <span class="pin" style="left: 40%; top: 12%">1</span><div class="bar"><span style="width:72%"></span></div>
        <table><tr><td>—</td><td>▲ 42%</td></tr></table><p>A → B</p></body></html>"""
        self.assertEqual(deck_tools.audit_design(clean), [])

    def test_shipped_templates_and_examples_are_clean(self) -> None:
        for html_path in [*pdfdesign.TEMPLATE_DIR.glob("*.html"), *(REPO_ROOT / "examples").glob("*.html")]:
            self.assertEqual(deck_tools.audit_design(html_path.read_text(encoding="utf-8")), [], html_path.name)

    def test_bundle_avoids_decorative_patterns(self) -> None:
        bundle = re.sub(r"/\*.*?\*/", "", pdfdesign.build_css_bundle(), flags=re.DOTALL)
        for pattern in ("decimal-leading-zero", "999px", "radial-gradient", "text-transform: uppercase", "box-shadow: 0 "):
            self.assertNotIn(pattern, bundle)
        self.assertIsNone(re.search(r"border-left:\s*[2-9](?:\.\d+)?pt solid var\(--accent\)", bundle))
        # 원형은 Figure 위치 핀(.pin, .callouts 번호)에만 허용한다.
        circle_rules = [rule for rule in re.findall(r"([^{}]+)\{[^}]*border-radius:\s*50%", bundle)]
        self.assertTrue(all(".pin" in rule or ".callouts" in rule for rule in circle_rules), circle_rules)
        self.assertIn(':root,\nhtml[data-palette="graphite"]', bundle, "기본 색조는 graphite")


class DocxWriterTest(unittest.TestCase):
    def test_docx_parts_are_well_formed(self) -> None:
        with tempfile.TemporaryDirectory(prefix="pdf-design-docx-") as work_dir:
            image_path = Path(work_dir) / "thumb.png"
            image_path.write_bytes(TINY_PNG)
            builder = DocxBuilder(title="테스트 & <대본>")
            builder.heading("제목", level=0)
            builder.heading("슬라이드 1", level=1)
            builder.paragraph([("굵게", True), (" 보통\x01", False)])
            builder.bullet("글머리표", level=1)
            builder.paragraph("대본 <문단>", style="Talk")
            builder.table(["#", "제목"], [["1", "A & B"]], [1.0, 5.0])
            builder.image(image_path, 5.0)
            builder.page_break()
            docx_path = builder.save(Path(work_dir) / "out.docx")

            with zipfile.ZipFile(docx_path) as archive:
                self.assertIsNone(archive.testzip())
                names = set(archive.namelist())
                for part in ("[Content_Types].xml", "word/document.xml", "word/styles.xml", "word/media/image1.png", "docProps/core.xml"):
                    self.assertIn(part, names)
                for part in ("word/document.xml", "word/styles.xml", "word/_rels/document.xml.rels", "docProps/core.xml"):
                    minidom.parseString(archive.read(part))  # 잘못된 XML이면 예외
                document_xml = archive.read("word/document.xml").decode("utf-8")
            self.assertNotIn("\x01", document_xml, "제어 문자는 제거돼야 한다")
            self.assertIn("대본 &lt;문단&gt;", document_xml)

            try:
                import docx  # python-docx 가 있으면 실제로 열리는지도 확인
            except ImportError:
                return
            document = docx.Document(str(docx_path))
            self.assertEqual(len(document.tables), 1)
            self.assertEqual(len(document.inline_shapes), 1)
            self.assertIn("Heading 1", {paragraph.style.name for paragraph in document.paragraphs})


BROKEN_DECK = """<!doctype html>
<html lang="en" data-style="academic" data-palette="graphite"><head><meta charset="utf-8"><title>Broken</title>
<link rel="stylesheet" href="pdf-design.css"><style>@page { size: 338.67mm 190.5mm; margin: 0; }</style></head>
<body class="deck">
<section class="slide">
  <header class="slide-head"><div class="slide-title">Overflowing body</div></header>
  <div class="slide-body l-points"><ol class="points">
    <li><strong>1</strong>x</li><li><strong>2</strong>x</li><li><strong>3</strong>x</li><li><strong>4</strong>x</li>
    <li><strong>5</strong>x</li><li><strong>6</strong>x</li><li><strong>7</strong>x</li><li><strong>8</strong>x</li>
  </ol></div>
  <footer class="slide-foot"><span class="source">Source</span><span class="page"></span></footer>
  <aside class="notes"><div class="talk"><p>대본</p></div></aside>
</section>
<section class="slide">
  <header class="slide-head"><div class="slide-title">Tiny text and missing image</div></header>
  <div class="slide-body"><p style="font-size: 8pt">too small</p><img src="missing.png" style="width: 40mm; height: 20mm"></div>
  <footer class="slide-foot"><span class="source"></span><span class="page"></span></footer>
</section>
</body></html>
"""


@unittest.skipUnless(HAS_RENDERER, "Chrome 또는 Poppler 가 없어 렌더링 테스트를 건너뜀")
class RenderSmokeTest(unittest.TestCase):
    def test_templates_render_expected_pages(self) -> None:
        expected_pages = {"onepager": 1, "deck": 3, "talk": 9}
        with tempfile.TemporaryDirectory(prefix="pdf-design-render-") as work_dir:
            for template, pages in expected_pages.items():
                doc = Path(work_dir) / f"{template}.html"
                self.assertEqual(run_cli("init", template, str(doc), "--style", "tech")[0], 0)
                exit_code, _, stderr = run_cli("render", str(doc), "--expect-pages", str(pages))
                self.assertEqual(exit_code, 0, stderr)
                self.assertTrue(doc.with_suffix(".pdf").exists())

    def test_check_passes_templates(self) -> None:
        with tempfile.TemporaryDirectory(prefix="pdf-design-check-") as work_dir:
            for template in ("talk", "deck", "onepager"):
                doc = Path(work_dir) / f"{template}.html"
                run_cli("init", template, str(doc))
                exit_code, stdout, stderr = run_cli("check", str(doc))
                self.assertEqual(exit_code, 0, stdout + stderr)
                self.assertIn("오류 0건", stdout)
                self.assertNotIn(deck_tools.AI_TELL_PREFIX, stdout)

    def test_check_detects_defects(self) -> None:
        with tempfile.TemporaryDirectory(prefix="pdf-design-check-") as work_dir:
            doc = Path(work_dir) / "broken.html"
            doc.write_text(BROKEN_DECK, encoding="utf-8")
            exit_code, stdout, _ = run_cli("check", str(doc))
            self.assertEqual(exit_code, 1)
            self.assertIn("[slide 1] 본문이 영역을 넘침", stdout)
            self.assertIn("핵심 포인트가 8개", stdout)
            self.assertIn("[slide 2] 글자가 너무 작음(8.0pt", stdout)
            self.assertIn("이미지를 불러오지 못함: missing.png", stdout)
            self.assertIn("[slide 2] 발표자 노트", stdout)
            self.assertFalse((Path(work_dir) / ".broken.check.html").exists(), "검사용 사본은 지워져야 한다")

    def test_script_builds_docx_with_thumbnails(self) -> None:
        with tempfile.TemporaryDirectory(prefix="pdf-design-script-") as work_dir:
            doc = Path(work_dir) / "talk.html"
            run_cli("init", "talk", str(doc))
            exit_code, stdout, stderr = run_cli("script", str(doc), "--minutes", "20")
            self.assertEqual(exit_code, 0, stderr)
            docx_path = Path(work_dir) / "talk_script.docx"
            self.assertTrue(docx_path.exists())
            self.assertIn("썸네일 포함", stdout)
            self.assertIn("목표 20분", stdout)
            with zipfile.ZipFile(docx_path) as archive:
                images = [name for name in archive.namelist() if name.startswith("word/media/")]
            self.assertEqual(len(images), 9, "슬라이드 수만큼 썸네일이 들어가야 한다")


if __name__ == "__main__":
    unittest.main()
